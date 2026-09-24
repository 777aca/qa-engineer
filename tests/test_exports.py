import json

import pytest
from openpyxl import load_workbook
from data_contract import DataError, load_document, normalize_document
from cases_to_xlsx import build_workbook
from cases_to_xmind import build_case_node, write_xmind
from bugs_to_xlsx import build_workbook as bug_workbook


def case():
    return {"id": "TC-1", "title": "=1+1", "module": "测试", "priority": "P1",
            "precondition": ["已登录"], "steps": ["执行操作"], "expected": ["结果正确"],
            "level": "L2", "platform": "mobile", "linked_bug": ["BUG-1"], "result": "Blocked",
            "actual": "缺少账号", "evidence": [{"type": "assertion", "content": "检查未执行"}]}


def test_legacy_fields_and_execution_metadata_survive_export(tmp_path):
    row = case()
    wb = build_workbook({"project": "项目:测试/资料", "cases": [row]}, "standard")
    output = tmp_path / "cases.xlsx"
    wb.save(output)
    sheet = load_workbook(output).active
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values["前置条件"] == "已登录"
    assert values["平台"] == "mobile"
    assert values["关联 Bug"] == "BUG-1"
    assert values["结果"] == "Blocked"
    assert values["用例标题"] == "=1+1"
    assert sheet.cell(2, 5).data_type == "s"
    xmind = json.dumps(build_case_node(row), ensure_ascii=False)
    assert all(value in xmind for value in ["已登录", "mobile", "BUG-1", "Blocked", "缺少账号"])
    assert "precondition" in row  # 不修改调用者输入


@pytest.mark.parametrize("change", [
    {"priority": "P9"}, {"platform": "desktop"}, {"result": "SUCCESS"},
    {"steps": []}, {"expected": {}}, {"preconditions": ["另一个值"]},
    {"platfrom": "pc"}, {"evidence": ["截图"]}, {"duration_ms": -1},
])
def test_invalid_cases_are_rejected(change):
    row = case()
    row.update(change)
    with pytest.raises(DataError):
        normalize_document({"cases": [row]})


def test_duplicate_ids_and_missing_collection_are_rejected():
    for data in ({}, {"cases": [case(), case()]}):
        with pytest.raises(DataError):
            normalize_document(data)


def test_existing_samples_and_zentao_still_work(tmp_path):
    data = load_document("samples/sample_cases.yaml")
    wb = build_workbook(data, "zentao")
    assert wb.active.max_column == 9
    assert wb.active.cell(2, 4).value
    write_xmind(data, tmp_path / "sample.xmind")
    assert (tmp_path / "sample.xmind").is_file()
    assert bug_workbook(load_document("samples/sample_bugs.yaml")).active.max_row > 4


def test_unconfirmed_bugs_are_default_and_confirmed_require_evidence():
    row = {"id": "B1", "title": "候选", "module": "页面", "severity": "S3", "priority": "P2"}
    normalized = normalize_document({"bugs": [row]}, "bugs")
    assert normalized["bugs"][0]["status"] == "待验证"
    with pytest.raises(DataError):
        normalize_document({"bugs": [{**row, "status": "已确认"}]}, "bugs")
